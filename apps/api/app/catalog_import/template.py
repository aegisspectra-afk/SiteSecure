"""Official SITE SECURE catalog import Excel template (no supplier data)."""

from __future__ import annotations

import io

from openpyxl import Workbook


def build_import_template_bytes() -> bytes:
    wb = Workbook()

    # Cameras
    ws = wb.active
    ws.title = "מצלמות"
    cam_headers = [
        "sku",
        "manufacturer",
        "model",
        "name",
        "description",
        "unit",
        "cost",
        "list_price",
        "resolution_mp",
        "environment",
        "form_factor",
        "poe",
        "max_power_w",
        "codec",
        "fps",
        "onvif",
        "lens_mm",
    ]
    ws.append(cam_headers)
    ws.append(
        [
            "CAM-EXAMPLE-4MP",
            "",
            "MODEL-4MP",
            "מצלמת דוגמה 4MP",
            "",
            "unit",
            "",
            "",
            "4",
            "outdoor",
            "bullet",
            "true",
            "",
            "H.265",
            "",
            "",
            "2.8",
        ]
    )

    nvr = wb.create_sheet("מקליטים")
    nvr.append(
        [
            "sku",
            "manufacturer",
            "model",
            "name",
            "description",
            "unit",
            "cost",
            "list_price",
            "channels",
            "poe_ports",
            "poe_budget_w",
            "drive_bays",
            "max_hdd_tb",
            "max_incoming_bandwidth_mbps",
            "codecs",
        ]
    )
    nvr.append(
        [
            "NVR-EXAMPLE-16",
            "",
            "NVR-16",
            "מקליט דוגמה 16 ערוצים",
            "",
            "unit",
            "",
            "",
            "16",
            "0",
            "",
            "2",
            "",
            "160",
            "",
        ]
    )

    hdd = wb.create_sheet("כוננים")
    hdd.append(
        [
            "sku",
            "manufacturer",
            "model",
            "name",
            "description",
            "unit",
            "cost",
            "list_price",
            "capacity_tb",
            "surveillance_grade",
        ]
    )
    hdd.append(
        [
            "HDD-EXAMPLE-8TB",
            "",
            "HDD-8TB",
            "דיסק דוגמה 8TB",
            "",
            "unit",
            "",
            "",
            "8",
            "true",
        ]
    )

    sw = wb.create_sheet("מתגים")
    sw.append(
        [
            "sku",
            "manufacturer",
            "model",
            "name",
            "description",
            "unit",
            "cost",
            "list_price",
            "ports",
            "poe_ports",
            "poe_budget_w",
            "port_speed_mbps",
            "uplink_speed_mbps",
        ]
    )
    sw.append(
        [
            "SW-EXAMPLE-16POE",
            "",
            "SW-16-POE",
            "מתג PoE דוגמה",
            "",
            "unit",
            "",
            "",
            "16",
            "16",
            "180",
            "1000",
            "1000",
        ]
    )

    other = wb.create_sheet("כללי")
    other.append(
        [
            "sku",
            "manufacturer",
            "model",
            "name",
            "description",
            "unit",
            "cost",
            "list_price",
        ]
    )

    # Instructions sheet
    info = wb.create_sheet("הוראות", 0)
    info.append(["תבנית ייבוא קטלוג — SITE SECURE"])
    info.append([])
    info.append(["1. מלאו גיליון לפי סוג המוצר (מצלמות / מקליטים / כוננים / מתגים / כללי)."])
    info.append(["2. תאים טכניים ריקים נשארים UNKNOWN — אל תמציאו ערכים."])
    info.append(["3. cost = עלות רכישה / מחיר מתקין; list_price = מחיר מכירה ללקוח."])
    info.append(["4. ניתן גם לייבא מחירון ספק דרך אשף המיפוי — התבנית אינה חובה."])
    info.append(["5. מחקו את שורות הדוגמה לפני ייבוא לייצור."])
    info.append(["6. משתמשי Google Sheets: הורידו כ-XLSX או CSV והעלו ל-SITE SECURE."])

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
