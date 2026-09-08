"""Unit tests for catalog import V1 — synthetic fixtures only (no supplier files)."""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from app.catalog_import.engine import (
    build_row_product,
    classify_duplicates,
    collect_candidates,
    readiness_from_products,
    summarize_candidates,
)
from app.catalog_import.mapping import suggest_column_map, suggest_field_for_header, suggest_sheet_include
from app.catalog_import.normalize import normalize_mapped_value
from app.catalog_import.parse import detect_header_row, parse_csv_bytes, parse_upload, parse_xlsx_bytes
from app.catalog_import.session import ImportSessionStore
from app.catalog_import.template import build_import_template_bytes


def _xlsx_bytes() -> bytes:
    wb = Workbook()
    menu = wb.active
    menu.title = "תפריט"
    menu.append(["מחירון דמו"])

    cam = wb.create_sheet(" IPC מצלמות רשת")
    cam.append(["חזרה לתפריט"])
    cam.append(["מק\"ט", "דגם", "מחיר מתקין", "תאור מקוצר", "רזולוציה", "סוג מצלמה", "גודל עדשה"])
    cam.append([None, "Series A"])  # banner
    cam.append(["32323212401", "IPC2124LB", 227, "מצלמת צינור 4MP", "4MP", "צינור", "2.8mm"])
    cam.append(["CAM-STR", "IPC-DOME", 300, "כיפה", "8MP", "כיפה", "2.8-12mm"])

    nvr = wb.create_sheet(" NVR מערכות הקלטה")
    nvr.append(["חזרה"])
    nvr.append(["מק\"ט", "דגם", "מחיר מתקין", "תאור מקוצר", "ערוצים", "כמות כוננים קשיחים", "רוחב פס נכנס", "יציאות POE"])
    nvr.append(["NVR-16", "NVR301-16", 469, "NVR 16", 16, 1, "80Mbps", "ללא"])

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def test_normalize_resolution_and_lens_and_mbps():
    assert normalize_mapped_value("attributes.resolution_mp", "4MP")["value"] == 4.0
    assert normalize_mapped_value("attributes.resolution_mp", "8MP")["value"] == 8.0
    assert normalize_mapped_value("attributes.lens_mm", "2.8mm")["value"] == 2.8
    assert normalize_mapped_value("attributes.lens_mm", "2.8-12mm")["value"] is None
    assert normalize_mapped_value("attributes.max_incoming_bandwidth_mbps", "80Mbps")["value"] == 80.0
    assert normalize_mapped_value("attributes.form_factor", "צינור")["value"] == "bullet"
    assert normalize_mapped_value("attributes.max_power_w", None)["value"] is None
    assert normalize_mapped_value("attributes.max_power_w", "")["provenance"] == "UNKNOWN"


def test_hebrew_header_auto_map():
    assert suggest_field_for_header('מק"ט') == "sku"
    assert suggest_field_for_header("דגם") == "model"
    assert suggest_field_for_header("מחיר מתקין") == "cost"
    assert suggest_field_for_header("רזולוציה") == "attributes.resolution_mp"
    assert suggest_field_for_header("ערוצים") == "attributes.channels"
    assert suggest_field_for_header("resolution_mp") == "attributes.resolution_mp"
    assert suggest_field_for_header("תמונה") is None


def test_parse_xlsx_multi_sheet_header_not_row1():
    parsed = parse_xlsx_bytes(_xlsx_bytes(), filename="demo.xlsx")
    assert parsed["format"] == "xlsx"
    assert len(parsed["sheets"]) == 3
    menu = parsed["sheets"][0]
    assert menu["suggested_include"] is False
    cam = parsed["sheets"][1]
    assert cam["header_row"] == 2
    assert cam["suggested_map"].get("0") == "sku"
    assert any(r and r[0] == "32323212401" for r in cam["rows"])


def test_parse_csv_and_reject_bad_type():
    csv_data = "sku,name,cost\nA1,Cam,10\n".encode("utf-8")
    parsed = parse_csv_bytes(csv_data, filename="a.csv")
    assert parsed["sheets"][0]["header_row"] == 1
    with pytest.raises(ValueError, match="UNSUPPORTED_TYPE"):
        parse_upload(b"not-a-file", filename="x.xls")


def test_numeric_and_string_sku_rows():
    product = build_row_product(
        ["32323212401", "IPC", 227, "name", "4MP", "צינור", "2.8mm"],
        column_map={
            "0": "sku",
            "1": "model",
            "2": "cost",
            "3": "name",
            "4": "attributes.resolution_mp",
            "5": "attributes.form_factor",
            "6": "attributes.lens_mm",
        },
        manufacturer_default="UNIVIEW",
        unit_default="unit",
        category_id="cat-1",
        category_key="cameras_ip",
        parent_key="video",
        can_set_cost=True,
    )
    assert product["status"] in {"ready", "warning"}
    assert product["product"]["sku"] == "32323212401"
    assert product["product"]["cost"] == 227.0
    assert product["product"]["list_price"] == 0.0
    assert "list_price_unset" in product["warnings"]
    assert product["product"]["attributes"]["resolution_mp"] == 4.0
    assert product["product"]["attributes"]["form_factor"] == "bullet"
    assert product["product"]["manufacturer"] == "UNIVIEW"
    # max_power unknown — not invented
    assert "max_power_w" not in product["product"]["attributes"]


def test_blocked_without_category_and_duplicate_policies():
    blocked = build_row_product(
        ["SKU1", "Name"],
        column_map={"0": "sku", "1": "name"},
        manufacturer_default=None,
        unit_default="unit",
        category_id=None,
        category_key=None,
        parent_key=None,
        can_set_cost=True,
    )
    assert blocked["status"] == "blocked"

    cands = [
        {
            "status": "ready",
            "product": {"sku": "A"},
        },
        {
            "status": "warning",
            "product": {"sku": "B"},
        },
    ]
    classify_duplicates(cands, {"A": "id-a"}, "skip")
    assert cands[0]["duplicate_action"] == "skip"
    assert cands[1]["duplicate_action"] == "create"
    classify_duplicates(cands, {"A": "id-a"}, "update")
    assert cands[0]["duplicate_action"] == "update"


def test_session_tenant_isolation():
    store = ImportSessionStore()
    s = store.create(
        workspace_id="ws-a",
        user_id="u1",
        filename="f.xlsx",
        format="xlsx",
        sheets=[],
    )
    assert store.get(s.id, workspace_id="ws-a") is not None
    assert store.get(s.id, workspace_id="ws-b") is None


def test_collect_candidates_skips_banners():
    parsed = parse_xlsx_bytes(_xlsx_bytes(), filename="demo.xlsx")
    cam = parsed["sheets"][1]
    cats = {
        "cat-cam": {"id": "cat-cam", "key": "cameras_ip", "parent_id": None},
    }
    cands = collect_candidates(
        parsed["sheets"],
        [
            {
                "sheet_index": cam["index"],
                "include": True,
                "header_row": 2,
                "category_id": "cat-cam",
                "manufacturer_default": "UNIVIEW",
                "unit_default": "unit",
                "column_map": cam["suggested_map"],
            }
        ],
        categories_by_id=cats,
        can_set_cost=True,
    )
    assert len(cands) == 2
    assert all(c["product"]["sku"] for c in cands)
    summary = summarize_candidates(cands)
    assert summary["detected"] == 2


def test_template_bytes_and_readiness():
    data = build_import_template_bytes()
    assert data[:2] == b"PK"
    readiness = readiness_from_products(
        [
            {"category_id": "c1", "attributes": {"resolution_mp": 4}},
            {"category_id": "c2", "attributes": {"channels": 16}},
            {"category_id": "c3", "attributes": {"capacity_tb": 8}},
        ],
        {
            "c1": {"key": "cameras_ip"},
            "c2": {"key": "nvr"},
            "c3": {"key": "hdd_recorders"},
        },
    )
    assert readiness["ready_for_core"] is True
    assert readiness["camera_structured"] == 1


def test_detect_header_row_prefers_product_headers():
    rows = [
        ["חזרה לתפריט"],
        ["מק\"ט", "דגם", "מחיר מתקין", "רזולוציה"],
        ["1", "M", 10, "4MP"],
    ]
    idx, conf = detect_header_row(rows)
    assert idx == 2
    assert conf > 0.3


def test_file_too_large():
    with pytest.raises(ValueError, match="FILE_TOO_LARGE"):
        parse_upload(b"x" * (12 * 1024 * 1024 + 1), filename="big.xlsx")


def test_importer_does_not_infer_environment_from_ip67_or_prose():
    from app.catalog_import.normalize import normalize_environment, normalize_mapped_value

    assert normalize_environment("IP67") == (None, "UNKNOWN")
    assert normalize_environment("waterproof IP67 enclosure") == (None, "UNKNOWN")
    assert normalize_environment("עמיד למים IP67") == (None, "UNKNOWN")
    # Explicit indoor/outdoor tokens still normalize — not IP67-derived
    assert normalize_environment("חוץ") == ("outdoor", "NORMALIZED")
    assert normalize_mapped_value("attributes.environment", "IP67")["value"] is None
    assert normalize_mapped_value("attributes.environment", "")["value"] is None
