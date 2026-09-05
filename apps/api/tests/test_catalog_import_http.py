"""HTTP + live catalog import recovery tests (Task 15B). Disposable fixtures only."""

from __future__ import annotations

import io
import uuid
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.config import get_settings
from app.main import app
from test_tenant_isolation import VIEWER_ID, _ensure_member, _password_grant, _rpc

pytestmark = pytest.mark.live

API_ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="module")
def settings():
    try:
        get_settings.cache_clear()
        return get_settings()
    except Exception:
        pytest.skip("SUPABASE_URL / SUPABASE_ANON_KEY not configured")


@pytest.fixture(scope="module")
def api():
    get_settings.cache_clear()
    return TestClient(app)


@pytest.fixture(scope="module")
def tenants(settings):
    password = "Test-Pass-2026!"
    token_a = _password_grant(settings, "ss.phase3.a@sitesecure.test", password)
    token_b = _password_grant(settings, "ss.phase3.b@sitesecure.test", password)
    ws_a = _rpc(settings, token_a, "create_workspace", {"p_name": f"Import A {uuid.uuid4().hex[:6]}", "p_plan_key": "solo"})
    ws_b = _rpc(settings, token_b, "create_workspace", {"p_name": f"Import B {uuid.uuid4().hex[:6]}", "p_plan_key": "solo"})
    assert ws_a.status_code == 200, ws_a.text
    assert ws_b.status_code == 200, ws_b.text
    return {
        "token_a": token_a,
        "token_b": token_b,
        "ws_a": ws_a.json(),
        "ws_b": ws_b.json(),
        "password": password,
    }


def _auth(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


def _xlsx_bytes() -> bytes:
    wb = Workbook()
    menu = wb.active
    menu.title = "תפריט"
    menu.append(["מחירון דמו"])
    cam = wb.create_sheet("IPC מצלמות")
    cam.append(["חזרה"])
    cam.append(["מק\"ט", "דגם", "מחיר מתקין", "תאור מקוצר", "רזולוציה", "סוג מצלמה", "גודל עדשה"])
    cam.append([None, "Series Banner"])
    cam.append([f"NUM-{uuid.uuid4().hex[:6]}", "IPC-BULLET", 227, "מצלמת צינור", "4MP", "צינור", "2.8mm"])
    cam.append([f"STR-{uuid.uuid4().hex[:6]}", "IPC-DOME", 300, "כיפה", "8MP", "כיפה", "2.8mm"])
    nvr = wb.create_sheet("NVR")
    nvr.append(["חזרה"])
    nvr.append(["מק\"ט", "דגם", "מחיר מתקין", "תאור מקוצר", "ערוצים", "כמות כוננים קשיחים", "רוחב פס נכנס", "יציאות POE"])
    nvr.append([f"NVR-{uuid.uuid4().hex[:6]}", "NVR-16", 469, "NVR 16", 16, 2, "80Mbps", "ללא"])
    hdd = wb.create_sheet("HDD")
    hdd.append(["מק\"ט", "דגם", "מחיר מתקין", "תאור מקוצר", "capacity_tb", "surveillance_grade"])
    hdd.append([f"HDD-{uuid.uuid4().hex[:6]}", "HDD-8", 450, "דיסק 8TB", 8, "true"])
    sw = wb.create_sheet("Switch")
    sw.append(["מק\"ט", "דגם", "מחיר מתקין", "תאור מקוצר", "ports", "poe_ports", "poe_budget_w"])
    sw.append([f"SW-{uuid.uuid4().hex[:6]}", "SW-16", 400, "מתג PoE", 16, 16, 180])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _csv_bytes(sku: str) -> bytes:
    return (
        "sku,name,cost,resolution_mp,form_factor\n"
        f"{sku},CSV Cam,150,4,bullet\n"
    ).encode("utf-8")


def _leaf(api: TestClient, token: str, ws: str, key: str) -> str:
    cats = api.get(f"/api/v1/workspaces/{ws}/catalog/categories", headers=_auth(token))
    assert cats.status_code == 200, cats.text
    hit = next((c for c in cats.json()["items"] if c.get("key") == key), None)
    assert hit, f"missing category {key}"
    return hit["id"]


def _sheet_cfg(meta: dict, category_id: str, manufacturer: str = "QABrand") -> dict:
    return {
        "sheet_index": meta["index"],
        "include": True,
        "header_row": meta["header_row"],
        "category_id": category_id,
        "manufacturer_default": manufacturer,
        "unit_default": "unit",
        "column_map": meta["suggested_map"],
    }


def test_owner_targets_template_parse_preview_commit_xlsx(api, tenants):
    ws = tenants["ws_a"]
    a = _auth(tenants["token_a"])

    targets = api.get(f"/api/v1/workspaces/{ws}/catalog/import/targets", headers=a)
    assert targets.status_code == 200, targets.text
    assert "fields" in targets.json()

    tpl = api.get(f"/api/v1/workspaces/{ws}/catalog/import/template", headers=a)
    assert tpl.status_code == 200, tpl.text
    assert tpl.content[:2] == b"PK"
    assert b"Uniview" not in tpl.content and b"Beres" not in tpl.content

    files = {"file": ("disp.xlsx", _xlsx_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    parsed = api.post(f"/api/v1/workspaces/{ws}/catalog/import/parse", headers=a, files=files)
    assert parsed.status_code == 200, parsed.text
    body = parsed.json()
    assert body["sheet_count"] >= 4
    sid = body["session_id"]
    sheets = {s["name"]: s for s in body["sheets"]}
    cam = next(s for s in body["sheets"] if "IPC" in s["name"] or "מצלמ" in s["name"])
    assert cam["header_row"] and cam["header_row"] > 1
    assert cam["suggested_map"].get("0") == "sku"
    assert "cost" in cam["suggested_map"].values()

    cam_id = _leaf(api, tenants["token_a"], ws, "cameras_ip")
    nvr_id = _leaf(api, tenants["token_a"], ws, "nvr")
    hdd_id = _leaf(api, tenants["token_a"], ws, "hdd_recorders")
    sw_id = _leaf(api, tenants["token_a"], ws, "switch")

    cfgs = []
    for s in body["sheets"]:
        name = s["name"]
        if not s["suggested_include"] and "תפריט" in name:
            continue
        if "IPC" in name or "מצלמ" in name:
            cfgs.append(_sheet_cfg(s, cam_id))
        elif "NVR" in name:
            cfgs.append(_sheet_cfg(s, nvr_id))
        elif "HDD" in name:
            # exact attribute headers
            m = dict(s["suggested_map"])
            cfgs.append({**_sheet_cfg(s, hdd_id), "column_map": m})
        elif "Switch" in name:
            cfgs.append(_sheet_cfg(s, sw_id))

    preview = api.post(
        f"/api/v1/workspaces/{ws}/catalog/import/preview",
        headers=a,
        json={"session_id": sid, "duplicate_policy": "skip", "sheets": cfgs},
    )
    assert preview.status_code == 200, preview.text
    summary = preview.json()["summary"]
    assert summary["detected"] >= 4
    assert summary["blocked"] == 0
    sample = preview.json()["sample_rows"][0]
    assert sample["product"]["list_price"] == 0
    assert sample["product"]["cost"] in {227, 300, 469, 450, 400} or sample["product"]["cost"] >= 0
    assert "list_price_unset" in (sample.get("warnings") or []) or sample["status"] in {"ready", "warning"}

    commit = api.post(
        f"/api/v1/workspaces/{ws}/catalog/import/commit",
        headers=a,
        json={"session_id": sid, "duplicate_policy": "skip", "sheets": cfgs, "confirm": True},
    )
    assert commit.status_code == 200, commit.text
    result = commit.json()
    assert result["imported"] >= 4
    assert result["readiness"]["camera_structured"] >= 1
    assert result["readiness"]["nvr_structured"] >= 1
    assert result["readiness"]["hdd_structured"] >= 1

    # session deleted
    again = api.post(
        f"/api/v1/workspaces/{ws}/catalog/import/preview",
        headers=a,
        json={"session_id": sid, "duplicate_policy": "skip", "sheets": cfgs},
    )
    assert again.status_code == 404

    # CCTV recommend consumes imported attributes
    rec = api.post(
        f"/api/v1/workspaces/{ws}/cctv/recommend",
        headers=a,
        json={
            "camera_count": 4,
            "retention_days": 14,
            "resolution_mp": 4,
            "poe_required": True,
            "architecture_intent": "prefer_external_switch",
        },
    )
    assert rec.status_code == 200, rec.text
    readiness = rec.json().get("catalog_readiness") or {}
    assert readiness.get("camera_structured", 0) >= 1
    assert readiness.get("nvr_structured", 0) >= 1
    assert readiness.get("hdd_structured", 0) >= 1


def test_csv_parse_preview_commit(api, tenants):
    ws = tenants["ws_a"]
    a = _auth(tenants["token_a"])
    sku = f"CSV-{uuid.uuid4().hex[:8]}"
    files = {"file": ("cams.csv", _csv_bytes(sku), "text/csv")}
    parsed = api.post(f"/api/v1/workspaces/{ws}/catalog/import/parse", headers=a, files=files)
    assert parsed.status_code == 200, parsed.text
    sheet = parsed.json()["sheets"][0]
    cam_id = _leaf(api, tenants["token_a"], ws, "cameras_ip")
    cfg = _sheet_cfg(sheet, cam_id)
    # ensure resolution maps
    if "3" not in cfg["column_map"]:
        cfg["column_map"]["3"] = "attributes.resolution_mp"
    if "4" not in cfg["column_map"]:
        cfg["column_map"]["4"] = "attributes.form_factor"
    body = {"session_id": parsed.json()["session_id"], "duplicate_policy": "skip", "sheets": [cfg]}
    assert api.post(f"/api/v1/workspaces/{ws}/catalog/import/preview", headers=a, json=body).status_code == 200
    commit = api.post(
        f"/api/v1/workspaces/{ws}/catalog/import/commit",
        headers=a,
        json={**body, "confirm": True},
    )
    assert commit.status_code == 200, commit.text
    assert commit.json()["imported"] == 1


def test_confirm_required_and_unauthorized(api, tenants, settings):
    ws = tenants["ws_a"]
    a = _auth(tenants["token_a"])
    files = {"file": ("a.csv", _csv_bytes(f"X-{uuid.uuid4().hex[:6]}"), "text/csv")}
    parsed = api.post(f"/api/v1/workspaces/{ws}/catalog/import/parse", headers=a, files=files)
    assert parsed.status_code == 200
    sheet = parsed.json()["sheets"][0]
    cam_id = _leaf(api, tenants["token_a"], ws, "cameras_ip")
    body = {
        "session_id": parsed.json()["session_id"],
        "duplicate_policy": "skip",
        "sheets": [_sheet_cfg(sheet, cam_id)],
        "confirm": False,
    }
    denied_confirm = api.post(f"/api/v1/workspaces/{ws}/catalog/import/commit", headers=a, json=body)
    assert denied_confirm.status_code == 400

    # viewer cannot commit
    _ensure_member(settings, tenants["token_a"], tenants["ws_a"], VIEWER_ID, "viewer")
    try:
        viewer_token = _password_grant(settings, "ss.phase3.viewer@sitesecure.test", tenants["password"])
    except Exception:
        pytest.skip("viewer test user unavailable")
    v = _auth(viewer_token)
    no = api.post(f"/api/v1/workspaces/{ws}/catalog/import/commit", headers=v, json={**body, "confirm": True})
    assert no.status_code in {403, 404}


def test_tenant_isolation_session_and_commit(api, tenants):
    ws_a = tenants["ws_a"]
    ws_b = tenants["ws_b"]
    a = _auth(tenants["token_a"])
    b = _auth(tenants["token_b"])

    files = {"file": ("iso.csv", _csv_bytes(f"ISO-{uuid.uuid4().hex[:6]}"), "text/csv")}
    parsed = api.post(f"/api/v1/workspaces/{ws_a}/catalog/import/parse", headers=a, files=files)
    assert parsed.status_code == 200
    sid = parsed.json()["session_id"]
    sheet = parsed.json()["sheets"][0]
    cam_a = _leaf(api, tenants["token_a"], ws_a, "cameras_ip")
    body = {"session_id": sid, "duplicate_policy": "skip", "sheets": [_sheet_cfg(sheet, cam_a)]}

    # B cannot preview A's session (even with A's session id on B workspace)
    prev_b = api.post(f"/api/v1/workspaces/{ws_b}/catalog/import/preview", headers=b, json=body)
    assert prev_b.status_code == 404

    # B cannot commit A's session into B
    commit_b = api.post(
        f"/api/v1/workspaces/{ws_b}/catalog/import/commit",
        headers=b,
        json={**body, "confirm": True},
    )
    assert commit_b.status_code == 404

    # A cannot commit into B with A's session
    commit_cross = api.post(
        f"/api/v1/workspaces/{ws_b}/catalog/import/commit",
        headers=a,
        json={**body, "confirm": True},
    )
    assert commit_cross.status_code in {403, 404}


def test_atomic_rollback_on_force_fail(api, tenants):
    ws = tenants["ws_a"]
    a = _auth(tenants["token_a"])
    sku = f"ROLL-{uuid.uuid4().hex[:8]}"
    files = {"file": ("roll.csv", _csv_bytes(sku), "text/csv")}
    parsed = api.post(f"/api/v1/workspaces/{ws}/catalog/import/parse", headers=a, files=files)
    assert parsed.status_code == 200
    sid = parsed.json()["session_id"]
    sheet = parsed.json()["sheets"][0]
    cam_id = _leaf(api, tenants["token_a"], ws, "cameras_ip")
    body = {
        "session_id": sid,
        "duplicate_policy": "skip",
        "sheets": [_sheet_cfg(sheet, cam_id)],
        "confirm": True,
        "test_force_fail": True,
    }
    fail = api.post(f"/api/v1/workspaces/{ws}/catalog/import/commit", headers=a, json=body)
    assert fail.status_code == 409, fail.text

    products = api.get(
        f"/api/v1/workspaces/{ws}/catalog/products",
        headers=a,
        params={"q": sku, "limit": 10, "include_inactive": "true"},
    )
    assert products.status_code == 200
    assert not any(p.get("sku") == sku for p in products.json().get("items", []))

    # retry without force works (session still alive)
    ok = api.post(
        f"/api/v1/workspaces/{ws}/catalog/import/commit",
        headers=a,
        json={**body, "test_force_fail": False},
    )
    assert ok.status_code == 200, ok.text
    assert ok.json()["imported"] == 1


def test_duplicate_skip_default(api, tenants):
    ws = tenants["ws_a"]
    a = _auth(tenants["token_a"])
    sku = f"DUP-{uuid.uuid4().hex[:8]}"
    files = {"file": ("dup.csv", _csv_bytes(sku), "text/csv")}
    parsed = api.post(f"/api/v1/workspaces/{ws}/catalog/import/parse", headers=a, files=files)
    sheet = parsed.json()["sheets"][0]
    cam_id = _leaf(api, tenants["token_a"], ws, "cameras_ip")
    body = {"session_id": parsed.json()["session_id"], "duplicate_policy": "skip", "sheets": [_sheet_cfg(sheet, cam_id)]}
    assert api.post(f"/api/v1/workspaces/{ws}/catalog/import/commit", headers=a, json={**body, "confirm": True}).status_code == 200

    parsed2 = api.post(f"/api/v1/workspaces/{ws}/catalog/import/parse", headers=a, files=files)
    body2 = {"session_id": parsed2.json()["session_id"], "duplicate_policy": "skip", "sheets": [_sheet_cfg(parsed2.json()["sheets"][0], cam_id)]}
    prev = api.post(f"/api/v1/workspaces/{ws}/catalog/import/preview", headers=a, json=body2)
    assert prev.status_code == 200
    assert prev.json()["summary"]["will_skip"] >= 1
    commit = api.post(f"/api/v1/workspaces/{ws}/catalog/import/commit", headers=a, json={**body2, "confirm": True})
    assert commit.status_code == 200
    assert commit.json()["imported"] == 0
    assert commit.json()["skipped"] >= 1
