"""FleetMaster CSV/XLSX import — parse + map + validate (no DB writes)."""

from __future__ import annotations

import csv
import io
import re
from typing import Any

from .fleet_logic import normalize_license_plate, validate_import_rows

MAX_FILE_BYTES = 8 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLS = 40
PREVIEW_ROWS = 10

_SAFE_NAME = re.compile(r"[^\w.\- ()\u0590-\u05FF]+", re.UNICODE)

# Hebrew + English header aliases → canonical field
COLUMN_ALIASES: dict[str, str] = {
    "license_plate": "license_plate",
    "plate": "license_plate",
    "לוחית": "license_plate",
    "לוחית רישוי": "license_plate",
    "מספר רישוי": "license_plate",
    "מספר רכב": "license_plate",
    "רכב": "license_plate",
    "mileage": "current_mileage",
    "current_mileage": "current_mileage",
    "קילומטראז": "current_mileage",
    "קילומטראז'": "current_mileage",
    'ק"מ': "current_mileage",
    "קמ": "current_mileage",
    "manufacturer": "manufacturer",
    "יצרן": "manufacturer",
    "model": "model",
    "דגם": "model",
    "status": "status",
    "סטטוס": "status",
    "מצב": "status",
    "vehicle_type": "vehicle_type",
    "סוג רכב": "vehicle_type",
    "סוג": "vehicle_type",
    "year": "year",
    "שנה": "year",
    "department": "department_key",
    "department_key": "department_key",
    "מחלקה": "department_key",
    "department_id": "department_id",
    "notes": "notes",
    "הערות": "notes",
    "test_expiry_date": "test_expiry_date",
    "תוקף טסט": "test_expiry_date",
    "טסט": "test_expiry_date",
    "insurance_expiry_date": "insurance_expiry_date",
    "תוקף ביטוח": "insurance_expiry_date",
    "ביטוח": "insurance_expiry_date",
    "insurance_provider": "insurance_provider",
    "חברת ביטוח": "insurance_provider",
    "insurance_policy_number": "insurance_policy_number",
    "פוליסה": "insurance_policy_number",
}


def sanitize_filename(name: str | None) -> str:
    base = (name or "upload").replace("\\", "/").split("/")[-1]
    base = _SAFE_NAME.sub("_", base).strip("._ ") or "upload"
    return base[:180]


def _norm_header(h: Any) -> str:
    s = str(h or "").strip().lower()
    s = s.replace("_", " ").replace("-", " ")
    s = re.sub(r"\s+", " ", s).strip()
    return s


def suggest_column_map(headers: list[Any]) -> dict[str, str]:
    """Map 0-based column index → canonical field."""
    out: dict[str, str] = {}
    used: set[str] = set()
    for i, h in enumerate(headers):
        key = _norm_header(h)
        # Direct + compacted keys
        candidates = [key, key.replace(" ", ""), str(h or "").strip()]
        field = None
        for c in candidates:
            if c in COLUMN_ALIASES:
                field = COLUMN_ALIASES[c]
                break
            # Partial Hebrew match
            for alias, target in COLUMN_ALIASES.items():
                if alias in c or c in alias:
                    field = target
                    break
            if field:
                break
        if field and field not in used:
            out[str(i)] = field
            used.add(field)
    return out


def _cell_value(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    if not s or s.startswith("="):
        return None
    return s


def _trim_row(vals: list[Any]) -> list[Any]:
    while vals and vals[-1] is None:
        vals.pop()
    return vals[:MAX_COLS]


def _score_header_row(vals: list[Any]) -> float:
    if not vals:
        return -1.0
    strs = sum(1 for v in vals if isinstance(v, str) and len(v.strip()) >= 2)
    if strs < 1:
        return -1.0
    joined = " ".join(str(v) for v in vals if isinstance(v, str)).lower()
    bonus = 0.0
    for token in ("רישוי", "לוחית", "plate", "קילומטר", "יצרן", "דגם", "mileage", "manufacturer"):
        if token in joined:
            bonus += 2.0
    return strs * 2.0 + bonus


def detect_header_row(rows: list[list[Any]], *, max_scan: int = 12) -> tuple[int | None, float]:
    best_i = None
    best_score = -1.0
    for i, row in enumerate(rows[:max_scan]):
        score = _score_header_row(row)
        if score > best_score:
            best_score = score
            best_i = i
    if best_i is None or best_score < 2:
        return None, 0.0
    return best_i + 1, min(1.0, best_score / 12.0)


def parse_csv_bytes(data: bytes, *, filename: str) -> dict[str, Any]:
    text = None
    for enc in ("utf-8-sig", "utf-8", "cp1255", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("ENCODING")
    reader = csv.reader(io.StringIO(text))
    rows: list[list[Any]] = []
    for i, row in enumerate(reader):
        if i >= MAX_ROWS:
            break
        rows.append(_trim_row([_cell_value(c) for c in row]))
    while rows and not any(v is not None for v in rows[-1]):
        rows.pop()
    header_row, conf = detect_header_row(rows)
    headers: list[Any] = list(rows[header_row - 1]) if header_row else []
    return {
        "filename": sanitize_filename(filename),
        "format": "csv",
        "header_row": header_row,
        "header_confidence": conf,
        "headers": headers,
        "suggested_map": suggest_column_map(headers) if headers else {},
        "preview_rows": rows[: PREVIEW_ROWS + (header_row or 1)],
        "rows": rows,
        "row_count": len(rows),
    }


def parse_xlsx_bytes(data: bytes, *, filename: str) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("OPENPYXL_MISSING") from exc

    bio = io.BytesIO(data)
    wb = load_workbook(bio, read_only=True, data_only=True, keep_links=False)
    try:
        ws = wb[wb.sheetnames[0]]
        rows: list[list[Any]] = []
        for ri, row in enumerate(ws.iter_rows(max_row=MAX_ROWS, max_col=MAX_COLS, values_only=True)):
            rows.append(_trim_row([_cell_value(c) for c in row]))
        while rows and not any(v is not None for v in rows[-1]):
            rows.pop()
    finally:
        wb.close()

    header_row, conf = detect_header_row(rows)
    headers: list[Any] = list(rows[header_row - 1]) if header_row else []
    return {
        "filename": sanitize_filename(filename),
        "format": "xlsx",
        "header_row": header_row,
        "header_confidence": conf,
        "headers": headers,
        "suggested_map": suggest_column_map(headers) if headers else {},
        "preview_rows": rows[: PREVIEW_ROWS + (header_row or 1)],
        "rows": rows,
        "row_count": len(rows),
    }


def parse_upload(data: bytes, *, filename: str, content_type: str | None = None) -> dict[str, Any]:
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("FILE_TOO_LARGE")
    if len(data) == 0:
        raise ValueError("EMPTY_FILE")
    name = sanitize_filename(filename).lower()
    ct = (content_type or "").lower()
    if name.endswith(".csv") or "csv" in ct:
        return parse_csv_bytes(data, filename=filename)
    if name.endswith((".xlsx", ".xlsm")) or "spreadsheet" in ct or "excel" in ct:
        return parse_xlsx_bytes(data, filename=filename)
    if name.endswith(".xls"):
        raise ValueError("UNSUPPORTED_TYPE")
    # sniff: try xlsx then csv
    if data[:2] == b"PK":
        return parse_xlsx_bytes(data, filename=filename)
    return parse_csv_bytes(data, filename=filename)


def apply_column_map(
    rows: list[list[Any]],
    *,
    header_row: int | None,
    column_map: dict[str, str],
) -> list[dict[str, Any]]:
    """Apply col-index → field map; skip header row."""
    start = header_row if header_row else 1
    mapped: list[dict[str, Any]] = []
    for row in rows[start:]:
        item: dict[str, Any] = {}
        empty = True
        for col_idx, field in column_map.items():
            try:
                i = int(col_idx)
            except (TypeError, ValueError):
                continue
            if i < 0 or i >= len(row):
                continue
            val = row[i]
            if val is not None and val != "":
                empty = False
            item[field] = val
        if empty:
            continue
        if "license_plate" not in item and "plate" in item:
            item["license_plate"] = item.pop("plate")
        mapped.append(item)
    return mapped


def preview_import(
    rows: list[dict[str, Any]],
    *,
    existing_plates: set[str] | frozenset[str] | None = None,
) -> dict[str, Any]:
    summary = validate_import_rows(rows, existing_plates=existing_plates)
    return {
        "summary": {
            "total": summary["total"],
            "valid": summary["valid"],
            "invalid": summary["invalid"],
            "duplicate_plates": summary["duplicate_plates"],
            "missing_plates": summary["missing_plates"],
        },
        "errors": summary["errors"],
        "warnings": summary["warnings"],
        "rows": summary["rows"][:50],
        "all_valid_rows": summary["rows"],
    }


def detect_file_duplicates(rows: list[dict[str, Any]]) -> list[str]:
    seen: set[str] = set()
    dups: list[str] = []
    for row in rows:
        plate = normalize_license_plate(row.get("license_plate") or row.get("plate"))
        if not plate:
            continue
        if plate in seen and plate not in dups:
            dups.append(plate)
        seen.add(plate)
    return dups
