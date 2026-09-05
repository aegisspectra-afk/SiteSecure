"""Parse XLSX/CSV into workspace-scoped import payloads (values only)."""

from __future__ import annotations

import csv
import io
import re
from typing import Any

from openpyxl import load_workbook

from .mapping import suggest_column_map, suggest_sheet_include

MAX_FILE_BYTES = 12 * 1024 * 1024
MAX_SHEETS = 40
MAX_ROWS_PER_SHEET = 5000
MAX_COLS = 60
PREVIEW_ROWS = 8

_SAFE_NAME = re.compile(r"[^\w.\- ()\u0590-\u05FF]+", re.UNICODE)


def sanitize_filename(name: str | None) -> str:
    base = (name or "upload").replace("\\", "/").split("/")[-1]
    base = _SAFE_NAME.sub("_", base).strip("._ ") or "upload"
    return base[:180]


def _cell_value(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    if not s:
        return None
    if s.startswith("="):
        # Unevaluated formula — treat as empty for import safety
        return None
    if s.upper() in {"#VALUE!", "#REF!", "#N/A", "#DIV/0!", "#NAME?", "#NULL!"}:
        return None
    return s


def _trim_row(vals: list[Any]) -> list[Any]:
    while vals and vals[-1] is None:
        vals.pop()
    return vals[:MAX_COLS]


def _score_header_row(vals: list[Any]) -> float:
    if not vals:
        return -1.0
    strs = sum(1 for v in vals if isinstance(v, str) and len(v.strip()) >= 2)
    nums = sum(1 for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool))
    if strs < 2:
        return -1.0
    joined = " ".join(str(v) for v in vals if isinstance(v, str)).lower()
    bonus = 0.0
    for token in ("מק", "דגם", "מחיר", "sku", "model", "price", "name", "תאור", "ערוצ", "רזולו"):
        if token in joined:
            bonus += 1.5
    return strs * 2.0 - nums + bonus


def detect_header_row(rows: list[list[Any]], *, max_scan: int = 15) -> tuple[int | None, float]:
    """Return 1-based row index and confidence 0..1."""
    best_i = None
    best_score = -1.0
    for i, row in enumerate(rows[:max_scan]):
        score = _score_header_row(row)
        if score > best_score:
            best_score = score
            best_i = i
    if best_i is None or best_score < 3:
        return None, 0.0
    # Normalize confidence
    conf = min(1.0, best_score / 12.0)
    return best_i + 1, conf


def parse_csv_bytes(data: bytes, *, filename: str) -> dict[str, Any]:
    text = None
    for enc in ("utf-8-sig", "utf-8", "cp1255", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("ENCODING")
    reader = csv.reader(io.StringIO(text))
    rows: list[list[Any]] = []
    for i, row in enumerate(reader):
        if i >= MAX_ROWS_PER_SHEET:
            break
        vals = _trim_row([_cell_value(c) for c in row])
        rows.append(vals)
    header_row, conf = detect_header_row(rows)
    headers: list[Any] = []
    if header_row:
        headers = list(rows[header_row - 1])
    return {
        "filename": sanitize_filename(filename),
        "format": "csv",
        "sheets": [
            {
                "index": 0,
                "name": "Sheet1",
                "row_count": len(rows),
                "suggested_include": True,
                "header_row": header_row,
                "header_confidence": conf,
                "headers": headers,
                "suggested_map": suggest_column_map(headers) if headers else {},
                "preview_rows": rows[: PREVIEW_ROWS + (header_row or 1)],
                "rows": rows,
            }
        ],
    }


def parse_xlsx_bytes(data: bytes, *, filename: str) -> dict[str, Any]:
    bio = io.BytesIO(data)
    # read_only + data_only: formulas become cached values or None — never execute macros
    wb = load_workbook(bio, read_only=True, data_only=True, keep_links=False)
    sheets_out: list[dict[str, Any]] = []
    for idx, name in enumerate(wb.sheetnames[:MAX_SHEETS]):
        ws = wb[name]
        rows: list[list[Any]] = []
        for ri, row in enumerate(ws.iter_rows(max_row=MAX_ROWS_PER_SHEET, max_col=MAX_COLS, values_only=True)):
            vals = _trim_row([_cell_value(c) for c in row])
            # Keep structural empties for header indexing; drop trailing all-empty at end later
            rows.append(vals)
        # Trim trailing fully empty rows
        while rows and not any(v is not None for v in rows[-1]):
            rows.pop()
        header_row, conf = detect_header_row(rows)
        headers: list[Any] = list(rows[header_row - 1]) if header_row else []
        sheets_out.append(
            {
                "index": idx,
                "name": name,
                "row_count": len(rows),
                "suggested_include": suggest_sheet_include(name) and len(rows) > 1,
                "header_row": header_row,
                "header_confidence": conf,
                "headers": headers,
                "suggested_map": suggest_column_map(headers) if headers else {},
                "preview_rows": rows[: PREVIEW_ROWS + (header_row or 1)],
                "rows": rows,
            }
        )
    wb.close()
    return {
        "filename": sanitize_filename(filename),
        "format": "xlsx",
        "sheets": sheets_out,
    }


def parse_upload(data: bytes, *, filename: str, content_type: str | None = None) -> dict[str, Any]:
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("FILE_TOO_LARGE")
    if len(data) == 0:
        raise ValueError("EMPTY_FILE")
    name = sanitize_filename(filename).lower()
    ct = (content_type or "").lower()
    is_csv = name.endswith(".csv") or "text/csv" in ct or "csv" in ct
    is_xlsx = (
        name.endswith(".xlsx")
        or "spreadsheetml" in ct
        or "officedocument.spreadsheetml" in ct
        or data[:2] == b"PK"  # zip/xlsx
    )
    if name.endswith(".xls") and not name.endswith(".xlsx"):
        raise ValueError("UNSUPPORTED_TYPE")
    if is_csv and not name.endswith(".xlsx"):
        return parse_csv_bytes(data, filename=filename)
    if is_xlsx:
        try:
            return parse_xlsx_bytes(data, filename=filename)
        except Exception as exc:  # noqa: BLE001 — surface as parse error
            raise ValueError("PARSE_FAILED") from exc
    raise ValueError("UNSUPPORTED_TYPE")
